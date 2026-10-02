"""Builds the Word (.docx) issue list."""
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt, RGBColor

SEVERITY_COLOR = {"error": RGBColor(0xC0, 0x1C, 0x28), "warning": RGBColor(0xB8, 0x6E, 0x00),
                  "check": RGBColor(0x1F, 0x5F, 0xAD)}


def timecode(seconds, fps):
    fps_whole = max(1, round(fps))
    frames = int(round(seconds * fps))
    secs, ff = divmod(frames, fps_whole)
    h, rem = divmod(secs, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}:{ff:02d}"


def build_docx(project_dir, name, analysis, review, issues, checklist, out_path):
    fps = analysis["info"]["fps"]
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.6))
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(f"Video QC Report: {name}", level=1)
    info = analysis.get("info", {})
    doc.add_paragraph(
        f"Duration: {timecode(info.get('duration', 0), fps)}    "
        f"Resolution: {info.get('width', '?')} x {info.get('height', '?')}    "
        f"FPS: {fps}"
    )
    if review.get("reviewer"):
        doc.add_paragraph(f"Reviewer: {review['reviewer']}")
    if review.get("summary"):
        doc.add_heading("Review Summary", level=2)
        doc.add_paragraph(str(review["summary"]))

    doc.add_heading("Checklist", level=2)
    checklist_values = review.get("checklist", {})
    for item in checklist:
        value = checklist_values.get(item, "Not marked")
        doc.add_paragraph(f"{item}: {value}", style="List Bullet")

    doc.add_heading("Transcript", level=2)
    transcript = analysis.get("transcript", [])
    if transcript:
        for segment in transcript:
            start = timecode(float(segment.get("start", 0)), fps)
            doc.add_paragraph(f"{start}  {segment.get('text', '')}")
    else:
        doc.add_paragraph("No narration transcript was generated.")

    kept = [i for i in issues if i["status"] != "dismissed"]
    doc.add_heading("Issues", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["Frame", "Timestamp", "Issue", "Inaccuracy / Finding", "Suggested Correction"]):
        cell.text = title
        cell.paragraphs[0].runs[0].bold = True
    widths = [Inches(1.0), Inches(1.1), Inches(1.9), Inches(3.8), Inches(2.2)]
    for issue in kept:
        cells = table.add_row().cells
        thumb = int(issue.get("thumb", issue.get("time", 0)))
        thumb_path = os.path.join(project_dir, "thumbs", f"{thumb}.jpg")
        if os.path.exists(thumb_path):
            cells[0].paragraphs[0].add_run().add_picture(thumb_path, width=Inches(0.9))
        else:
            cells[0].text = "-"
        tc = timecode(issue["time"], fps)
        if issue.get("end"):
            tc += f"\nto {timecode(issue['end'], fps)}"
        cells[1].text = tc
        cells[2].text = issue.get("title") or issue["category"]
        p = cells[3].paragraphs[0]
        category = p.add_run(f"{issue['category']}: ")
        category.bold = True
        category.font.color.rgb = SEVERITY_COLOR.get(issue.get("severity"), RGBColor(0, 0, 0))
        if issue.get("detail"):
            p.add_run(issue["detail"])
        else:
            p.add_run(issue.get("title", ""))
        cells[4].text = issue.get("note", "") or (issue.get("suggestions") or {}).get("correction", "")
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

    doc.save(out_path)
